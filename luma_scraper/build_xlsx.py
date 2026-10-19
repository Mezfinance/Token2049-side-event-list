#!/usr/bin/env python3
"""Step 3c: build styled XLSX files with the official TOKEN2049 Week banner.

Produces two workbooks from data/export/rows.json, styled to match the
official side-events Google Sheet (navy header row, light-blue date column):

  * TOKEN2049_2026_Luma_Organisers.xlsx  - one row per event
  * TOKEN2049_2026_Luma_Hosts.xlsx       - one row per host per event

Usage:
    python3 build_xlsx.py --rows data/export/rows.json --out-dir data/export
"""
from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
BANNER = HERE / "assets" / "token2049_week_banner.png"

NAVY = "073763"
LIGHT_BLUE = "6FA8DC"
PALE_BLUE = "CFE2F3"
GREY_LINE = "D9D9D9"
FONT = "Calibri"

HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=12)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name=FONT, bold=True, color="1A1A1A", size=16)
NOTE_FONT = Font(name=FONT, italic=True, color="595959", size=10)
BODY_FONT = Font(name=FONT, color="212529", size=11)
LINK_FONT = Font(name=FONT, color="0563C1", underline="single", size=11)
DATE_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
BAND_FILL = PatternFill("solid", fgColor="F3F6FA")
THIN = Side(style="thin", color=GREY_LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BANNER_ROWS = 3          # rows 1-3 hold the banner
TITLE_ROW = 4
NOTE_ROW = 5
HEADER_ROW = 6
FIRST_DATA_ROW = 7

# column widths per header name; anything not listed gets a default
WIDTHS = {
    "Date": 17, "Start": 8, "End": 8, "Event": 48, "Venue": 42, "Venue status": 22,
    "Luma link": 30, "Platform": 9,
    "Tags": 20, "Organising calendar": 30, "Calendar website": 30, "Calendar X": 28,
    "Calendar LinkedIn": 34, "Hosts": 40, "Host bios": 60, "Host websites": 36,
    "Host X handles": 36, "Host LinkedIn": 40, "Host Instagram": 32, "Host count": 7,
    "Guest count": 8, "City": 12, "Country": 12, "Event ID": 20,
    "Host name": 28, "Host username": 16, "Host bio": 60, "Host website": 32,
    "Host X": 30, "Host Instagram": 32, "Host TikTok": 14, "Host YouTube": 16, "Verified": 8,
}
LINK_COLUMNS = {
    "Luma link", "Calendar website", "Calendar X", "Calendar LinkedIn",
    "Host website", "Host X", "Host LinkedIn", "Host Instagram",
}
# "Host bios" (events view) stays unwrapped: with 9 hosts in one cell a wrapped row gets very tall
WRAP_COLUMNS = {"Event", "Venue", "Hosts", "Host bio", "Tags", "Organising calendar"}
CENTER_COLUMNS = {"Start", "End", "Platform", "Host count", "Guest count", "Verified"}


def add_banner(ws, n_cols: int) -> None:
    img = XLImage(str(BANNER))
    # banner is 2048x228; scale to ~110px tall and keep the aspect ratio
    scale = 110 / img.height
    img.height = int(img.height * scale)
    img.width = int(img.width * scale)
    for r in range(1, BANNER_ROWS + 1):
        ws.row_dimensions[r].height = 30
    ws.add_image(img, "A1")


def style_sheet(ws, columns: list[str], rows: list[list], title: str, note: str) -> None:
    n_cols = len(columns)
    add_banner(ws, n_cols)

    ws.cell(row=TITLE_ROW, column=1, value=title).font = TITLE_FONT
    ws.row_dimensions[TITLE_ROW].height = 26
    ws.cell(row=NOTE_ROW, column=1, value=note).font = NOTE_FONT
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=min(n_cols, 8))
    ws.merge_cells(start_row=NOTE_ROW, start_column=1, end_row=NOTE_ROW, end_column=min(n_cols, 8))

    for c, name in enumerate(columns, 1):
        cell = ws.cell(row=HEADER_ROW, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS.get(name, 18)
    ws.row_dimensions[HEADER_ROW].height = 30

    link_idx = {i for i, n in enumerate(columns) if n in LINK_COLUMNS}
    wrap_idx = {i for i, n in enumerate(columns) if n in WRAP_COLUMNS}
    center_idx = {i for i, n in enumerate(columns) if n in CENTER_COLUMNS}
    date_idx = columns.index("Date") if "Date" in columns else None

    for r_off, row in enumerate(rows):
        r = FIRST_DATA_ROW + r_off
        banded = r_off % 2 == 1
        for i, value in enumerate(row):
            cell = ws.cell(row=r, column=i + 1, value=value if value != "" else None)
            cell.border = BORDER
            if i in link_idx and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = LINK_FONT
            else:
                cell.font = BODY_FONT
            if i == date_idx:
                cell.fill = DATE_FILL
            elif banded:
                cell.fill = BAND_FILL
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=i in wrap_idx,
                horizontal="center" if i in center_idx else None,
            )

    last_row = FIRST_DATA_ROW + len(rows) - 1
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(n_cols)}{last_row}"
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY


def build(rows_path: Path, out_dir: Path, fetched: str) -> list[Path]:
    data = json.loads(rows_path.read_text())
    n_events = len(data["events"]["rows"])
    n_hosts = len(data["hosts"]["rows"])
    source = "Source: Luma TOKEN2049 calendar, https://luma.com/token2049?period=past, fetched " + fetched
    specs = [
        ("TOKEN2049_2026_Luma_Organisers.xlsx", "Luma Organisers", data["events"],
         f"TOKEN2049 Week Singapore 2026: side-event organisers ({n_events} events)",
         f"One row per event; several hosts in a cell are separated by '; '. {source}."),
        ("TOKEN2049_2026_Luma_Hosts.xlsx", "Luma Hosts", data["hosts"],
         f"TOKEN2049 Week Singapore 2026: side-event hosts ({n_hosts} host entries, {n_events} events)",
         f"One row per host per event; filter on Host name to see everything a person or organisation ran. {source}."),
    ]
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for filename, tab, table, title, note in specs:
        wb = Workbook()
        ws = wb.active
        ws.title = tab
        style_sheet(ws, table["columns"], table["rows"], title, note)
        path = out_dir / filename
        wb.save(path)
        written.append(path)
        print(f"wrote {path} ({len(table['rows'])} rows)")
    return written


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--rows", default="data/export/rows.json")
    ap.add_argument("--out-dir", default="data/export")
    ap.add_argument("--fetched", default=date.today().isoformat(), help="date the data was pulled")
    args = ap.parse_args()
    build(Path(args.rows), Path(args.out_dir), args.fetched)


if __name__ == "__main__":
    main()
